import copy
import io
import json
import os
import tempfile
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from uuid import uuid4

TEST_DIR = tempfile.TemporaryDirectory(prefix='gesto-test-')
os.environ['DATA_DIR'] = TEST_DIR.name
os.environ['ADMIN_PASSWORD'] = 'test-owner-password'
os.environ['APP_ENV'] = 'development'
os.environ.pop('SECRET_KEY', None)
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from server.app import app
from server.engine import validate_document, evaluate, FormError
from server.export import make_workbook
from server.seed import example_document, blank_document
from server.storage import connect, dumps
from server.google import process_jobs, sync_job, worker_lock, GoogleError


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.doc = example_document()

    def test_title_and_option_style_validation_is_backwards_compatible(self):
        self.assertEqual(validate_document(self.doc), [])
        self.doc['style'].update(titleText='#123456', optionText='#abcdef', titleSize='huge')
        self.assertEqual(validate_document(self.doc), [])
        self.doc['questions'][0]['style']['titleSize'] = 'arbitrary-css'
        self.assertTrue(any('Tamaño' in e for e in validate_document(self.doc)))
        self.doc['questions'][0]['style'] = {'optionText': 'red;display:none'}
        self.assertTrue(any('Color' in e for e in validate_document(self.doc)))

    def test_footer_is_optional_text_with_limit(self):
        self.doc['footer'] = ''
        self.assertEqual(validate_document(self.doc), [])
        self.doc['footer'] = 'Club Gesto'
        self.assertEqual(validate_document(self.doc), [])
        for footer in ['x' * 101, 42]:
            self.doc['footer'] = footer
            self.assertTrue(any('pie' in e for e in validate_document(self.doc)))

    def test_valid_example_two_branches_and_boundaries(self):
        self.assertEqual(validate_document(self.doc), [])
        result = evaluate(self.doc, {'interes': 'fuerza', 'entrenamiento': 'regular', 'nombre': 'Ana'})
        self.assertEqual(result['total'], 6)
        self.assertEqual(result['result']['id'], 'activo')
        result = evaluate(self.doc, {'interes': 'equilibrio', 'bienestar': 'pausa'})
        self.assertEqual(result['total'], 1)
        self.assertEqual([a['question_id'] for a in result['answers']], ['interes', 'bienestar', 'nombre'])

    def test_stale_branch_and_required_rejected(self):
        with self.assertRaises(FormError): evaluate(self.doc, {'interes': 'equilibrio', 'bienestar': 'pausa', 'entrenamiento': 'regular'})
        with self.assertRaises(FormError): evaluate(self.doc, {'interes': 'fuerza'})

    def test_cycles_dangling_destinations_even_unreachable(self):
        self.doc['questions'][2]['next'] = 'interes'
        self.assertTrue(any('bucle' in e for e in validate_document(self.doc)))
        self.doc['questions'][2]['next'] = 'deleted'
        self.assertTrue(any('Destino' in e for e in validate_document(self.doc)))

    def test_ranges_overlap_gap_default_and_open(self):
        self.doc['results'][0]['max'] = 5
        self.assertTrue(any('solapan' in e for e in validate_document(self.doc)))
        self.doc['results'][0]['max'] = 2
        self.assertTrue(any('Cubre' in e for e in validate_document(self.doc)))
        self.doc['results'].append({'id': 'fallback', 'title': 'Otro', 'default': True})
        self.assertEqual(validate_document(self.doc), [])

    def test_optional_default_end_negative_and_zero(self):
        q = self.doc['questions'][0]
        q['required'] = False
        q['next'] = 'end'
        result = evaluate(self.doc, {})
        self.assertEqual(result['total'], 0)
        q['options'][0]['points'] = -7
        q['options'][0]['next'] = 'end'
        self.assertEqual(evaluate(self.doc, {'interes':'fuerza'})['total'], -7)

    def test_statement_screen_has_no_answer_points_or_column(self):
        # A text screen between the first choice and the rest of the route, also used as a jump target.
        self.doc['questions'].insert(1, {'id': 'aviso', 'type': 'statement', 'title': 'Antes de seguir', 'description': 'Texto', 'button': 'Vamos', 'next': 'next', 'style': {'textSize': 'large'}})
        self.doc['questions'][0]['options'][0]['next'] = 'aviso'
        self.assertEqual(validate_document(self.doc), [])
        result = evaluate(self.doc, {'interes': 'fuerza', 'entrenamiento': 'regular'})
        self.assertEqual(result['total'], 6)
        self.assertNotIn('aviso', [a['question_id'] for a in result['answers']])
        with self.assertRaises(FormError):
            evaluate(self.doc, {'interes': 'fuerza', 'aviso': 'x', 'entrenamiento': 'regular'})
        submission = {'id': 's1', 'version': 1, 'completed': '2026-10-23T10:00:00+00:00', 'total': 6, 'result_title': 'R', 'answers': json.dumps(result['answers'])}
        headers = [c.value for c in load_workbook(io.BytesIO(make_workbook([submission], [self.doc]))).active[1]]
        self.assertFalse(any('[aviso]' in h for h in headers))
        self.doc['questions'][1]['button'] = 'x' * 101
        self.assertTrue(any('botón' in e for e in validate_document(self.doc)))

    def test_text_size_is_validated(self):
        self.doc['style']['textSize'] = 'xsmall'
        self.assertEqual(validate_document(self.doc), [])
        self.doc['questions'][0]['style']['textSize'] = '200%'
        self.assertTrue(any('texto' in e for e in validate_document(self.doc)))


class APITests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        response = self.client.post('/api/auth/login', json={'password': 'test-owner-password'})
        self.assertEqual(response.status_code, 200)
        self.fid = self.client.post('/api/forms', json={'name': 'Evaluación ágil'}).json()['id']
        self.doc = example_document()
        self.save(self.doc)

    def save(self, doc):
        current = self.client.get('/api/forms/' + self.fid).json()
        result = self.client.put('/api/forms/' + self.fid, json={'draft': doc, 'revision': current['revision']})
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    def publish(self):
        self.assertEqual(self.client.post(f'/api/forms/{self.fid}/publish').status_code, 200)
        return self.client.get('/api/public/' + self.fid).json()

    def submission(self, version=None, answers=None, sid=None):
        version = version or self.publish()
        return {'id': sid or str(uuid4()), 'version': version['version'], 'token': version['token'], 'answers': answers if answers is not None else {'interes':'fuerza','entrenamiento':'regular','nombre':'Lucía'}}

    def test_private_admin_export_and_csrf(self):
        client = TestClient(app)
        for path in ['/api/forms', f'/api/forms/{self.fid}', f'/api/forms/{self.fid}/submissions', f'/api/forms/{self.fid}/export', f'/api/forms/{self.fid}/integration']:
            self.assertEqual(client.get(path).status_code, 401)
        self.assertEqual(self.client.post('/api/forms', json={'name':'bad'}, headers={'Origin':'https://evil.example'}).status_code, 403)

    def test_version_pinning_draft_isolation_and_idempotency(self):
        public = self.publish()
        payload = self.submission(public)
        self.doc['questions'][0]['options'][0]['points'] = 100
        self.save(self.doc)
        self.assertEqual(self.client.get('/api/public/' + self.fid).json()['document']['questions'][0]['options'][0]['points'], 3)
        latest = self.publish()
        response = self.client.post(f'/api/public/{self.fid}/submit', json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['total'], 6)
        self.assertEqual(self.client.post(f'/api/public/{self.fid}/submit', json=payload).status_code, 200)
        self.assertEqual(len(self.client.get(f'/api/forms/{self.fid}/submissions').json()), 1)
        payload['answers']['nombre'] = 'Cambio'
        self.assertEqual(self.client.post(f'/api/public/{self.fid}/submit', json=payload).status_code, 409)
        new = self.client.post(f'/api/public/{self.fid}/submit', json=self.submission(latest)).json()
        self.assertEqual(new['total'], 103)

    def test_concurrent_double_click_has_one_submission(self):
        payload = self.submission()
        with ThreadPoolExecutor(max_workers=4) as pool:
            statuses = list(pool.map(lambda _: TestClient(app).post(f'/api/public/{self.fid}/submit', json=payload).status_code, range(4)))
        self.assertEqual(statuses, [200]*4)
        self.assertEqual(len(self.client.get(f'/api/forms/{self.fid}/submissions').json()), 1)

    def test_server_rejects_tampered_data(self):
        payload = self.submission()
        payload['answers']['bienestar'] = 'pausa'
        self.assertEqual(self.client.post(f'/api/public/{self.fid}/submit', json=payload).status_code, 422)
        del payload['answers']['bienestar']
        payload['total'] = 99999
        response = self.client.post(f'/api/public/{self.fid}/submit', json=payload)
        self.assertEqual(response.json()['total'], 6)
        payload['token'] = 'forged'
        self.assertEqual(self.client.post(f'/api/public/{self.fid}/submit', json=payload).status_code, 403)

    def test_publish_validation_and_revision_conflict(self):
        self.doc['questions'][0]['next'] = 'gone'
        self.save(self.doc)
        self.assertEqual(self.client.post(f'/api/forms/{self.fid}/publish').status_code, 422)
        self.assertEqual(self.client.put(f'/api/forms/{self.fid}', json={'draft':self.doc,'revision':1}).status_code, 409)

    def test_xlsx_historical_stable_columns_zero_and_formula_strings(self):
        doc = blank_document('Excel')
        doc['questions'] = [
            {'id':'zero','title':'Número','type':'number','required':True,'next':'next'},
            {'id':'text','title':'Texto','type':'short','required':False,'next':'end'},
            {'id':'skipped','title':'No mostrada','type':'short','required':False,'next':'end'}]
        self.save(doc)
        old = self.publish()
        formula = '=HYPERLINK("https://example.com")'
        payload = self.submission(old, {'zero':0, 'text':formula})
        self.assertEqual(self.client.post(f'/api/public/{self.fid}/submit', json=payload).status_code, 200)
        doc['questions'] = [{'id':'text','title':'Nombre cambiado','type':'short','required':False,'next':'end'}]
        self.save(doc)
        new = self.publish()
        self.client.post(f'/api/public/{self.fid}/submit', json=self.submission(new, {'text':'áéíóú Ñ @literal'}))
        export = self.client.get(f'/api/forms/{self.fid}/export')
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(io.BytesIO(export.content)).active
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:H3')
        self.assertEqual(sheet['F2'].value, 0)
        self.assertEqual(sheet['G2'].value, formula)
        self.assertEqual(sheet['G2'].data_type, 's')
        self.assertIsNone(sheet['H2'].value)
        self.assertEqual(sheet['G3'].value, 'áéíóú Ñ @literal')
        self.assertEqual(sheet['C2'].data_type, 'd')
        filtered = self.client.get(f'/api/forms/{self.fid}/export?start=2000-01-01&end=2000-01-02')
        self.assertEqual(load_workbook(io.BytesIO(filtered.content)).active.max_row, 1)

    def test_images_validate_content_and_persist(self):
        self.assertEqual(self.client.post('/api/uploads',files={'file':('evil.png',b'<script>','image/png')}).status_code,422)
        stream = io.BytesIO(); Image.new('RGB',(20,20),'green').save(stream,'PNG')
        response = self.client.post('/api/uploads',files={'file':('valid.png',stream.getvalue(),'image/png')})
        self.assertEqual(response.status_code,200)
        url = response.json()['url']
        self.assertEqual(TestClient(app).get(url).status_code,200)

    def test_gallery_lists_uploads_newest_first_and_is_private(self):
        urls = []
        for color in ['red', 'blue']:
            stream = io.BytesIO(); Image.new('RGB',(10,10),color).save(stream,'PNG')
            urls.append(self.client.post('/api/uploads',files={'file':('x.png',stream.getvalue(),'image/png')}).json()['url'])
            time.sleep(0.02)
        gallery = [item['url'] for item in self.client.get('/api/uploads').json()]
        self.assertEqual(gallery[:2], urls[::-1])
        self.assertIn(TestClient(app).get('/api/uploads').status_code, (401, 403))

    def test_duplicate_is_draft_only_and_delete_cascades(self):
        self.client.post(f'/api/public/{self.fid}/submit',json=self.submission())
        copy_id = self.client.post(f'/api/forms/{self.fid}/duplicate').json()['id']
        self.assertIsNone(self.client.get('/api/forms/'+copy_id).json()['published'])
        self.assertEqual(self.client.get('/api/public/'+copy_id).status_code,404)
        self.assertEqual(self.client.delete('/api/forms/'+self.fid).status_code,200)
        with connect() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM submissions WHERE form_id=?',(self.fid,)).fetchone()[0],0)

    def test_persistence_survives_new_connection(self):
        payload = self.submission()
        self.client.post(f'/api/public/{self.fid}/submit',json=payload)
        with connect() as db:
            self.assertEqual(db.execute('SELECT id FROM submissions WHERE id=?',(payload['id'],)).fetchone()['id'],payload['id'])

    def test_text_long_email_numeric_and_nonfinite_validation(self):
        doc = blank_document('Formatos')
        doc['questions'] = [
            {'id':'email','title':'Correo','type':'email','required':True,'next':'next'},
            {'id':'number','title':'Número','type':'number','required':True,'next':'next'},
            {'id':'long','title':'Comentarios','type':'long','required':True,'next':'end'}]
        self.save(doc)
        public=self.publish()
        for answers in [
            {'email':'mal@','number':0,'long':'Hola'},
            {'email':'a@b.com','number':True,'long':'Hola'},
            {'email':'a@b.com','number':None,'long':'Hola'},
            {'email':'a@b.com','number':0,'long':'  '},
            {'email':'a@b.com','number':0,'long':'control\x00'}]:
            self.assertEqual(self.client.post(f'/api/public/{self.fid}/submit',json=self.submission(public,answers)).status_code,422)
        good=self.client.post(f'/api/public/{self.fid}/submit',json=self.submission(public,{'email':'a@b.com','number':0,'long':'Línea 1\nLínea 2'}))
        self.assertEqual(good.status_code,200)
        self.assertEqual(self.client.post('/api/forms',content='{"name":NaN}',headers={'Content-Type':'application/json'}).status_code,400)


class FakeSheets:
    """Transport-only test double; never presented as a real Google verification."""
    def __init__(self):
        self.rows = {}; self.fail_after_write = False; self.calls = []
    def metadata(self, spreadsheet):
        return {'sheets':[{'properties':{'sheetId':0,'title':'Respuestas','gridProperties':{'rowCount':1000,'columnCount':50}}}]}
    def values(self, spreadsheet, cell_range):
        if cell_range.endswith('1:1'): return [self.rows[1]] if 1 in self.rows else []
        if cell_range.endswith('A:A'): return [[self.rows[i][0]] if i in self.rows else [] for i in range(1,max(self.rows,default=0)+1)]
        raise AssertionError(cell_range)
    def ensure_size(self,*args): pass
    def write(self,spreadsheet,range,rows):
        row=int(range.rsplit('A',1)[1]);self.rows[row]=rows[0];self.calls.append(row)
        if self.fail_after_write and row>1:
            self.fail_after_write=False;raise OSError('Lost success response')


class SyncTests(unittest.TestCase):
    setUp = APITests.setUp
    save = APITests.save
    publish = APITests.publish
    submission = APITests.submission
    def test_sync_retry_after_uncertain_success_and_original_destination(self):
        destination=str(uuid4())
        with connect() as db:
            db.execute('INSERT INTO destinations(id,spreadsheet,sheet_id,title) VALUES(?,?,0,?)',(destination,'document-test-'+destination,'Respuestas'))
            db.execute('INSERT INTO integrations VALUES(?,?,1)',(self.fid,destination))
        public=self.publish();payload=self.submission(public)
        self.client.post(f'/api/public/{self.fid}/submit',json=payload)
        with connect() as db:
            row=db.execute('SELECT * FROM jobs WHERE submission_id=?',(payload['id'],)).fetchone()
            self.assertEqual(row['status'],'pending')
            self.assertEqual(row['destination_id'],destination)
            db.execute('UPDATE integrations SET enabled=0 WHERE form_id=?',(self.fid,))
        fake=FakeSheets();fake.fail_after_write=True
        process_jobs(fake)
        with connect() as db:
            job=db.execute('SELECT * FROM jobs WHERE submission_id=?',(payload['id'],)).fetchone()
            self.assertEqual(job['status'],'error');self.assertEqual(job['attempts'],1)
            self.assertEqual(job['row_number'],2)
            self.assertEqual(job['destination_id'],destination)
            self.assertGreater(job['next_attempt'],0)
        self.assertEqual(self.client.post('/api/submissions/'+payload['id']+'/retry').status_code,200)
        process_jobs(fake)
        self.assertEqual(len(fake.rows),2)
        self.assertEqual(fake.rows[2][0],payload['id'])
        with connect() as db:
            self.assertEqual(db.execute('SELECT status FROM jobs WHERE submission_id=?',(payload['id'],)).fetchone()['status'],'synced')
        # A second version renames/reorders; old columns keep their mapping.
        self.doc['questions'][-1]['title']='Duplicated title'
        self.doc['questions'].append({'id':'new','title':'Duplicated title','type':'short','required':False,'next':'end'})
        self.doc['questions'][-2]['next']='next'
        self.save(self.doc);new=self.publish()
        with connect() as db: db.execute('UPDATE integrations SET enabled=1 WHERE form_id=?',(self.fid,))
        second=self.submission(new,{'interes':'equilibrio','bienestar':'pausa','nombre':'=literal','new':'+literal'})
        self.client.post(f'/api/public/{self.fid}/submit',json=second)
        process_jobs(fake)
        self.assertEqual(fake.rows[3][-2:],['=literal','+literal'])
        self.assertEqual(fake.rows[1][-2].split(' [')[0],'¿Cómo te llamas?')
        self.assertEqual(len(fake.rows[1]),10)

    def test_process_lock_excludes_other_worker(self):
        with worker_lock() as first:
            self.assertTrue(first)
            with worker_lock() as second:
                self.assertFalse(second)

    def test_google_uses_raw_literals_and_refreshes_token(self):
        from unittest.mock import patch
        import httpx
        from server.app import google
        google.store_tokens({'access_token':'test-only','refresh_token':'test-only-refresh','expires_in':-1})
        with patch('server.google.httpx.post',return_value=httpx.Response(200,json={'access_token':'renewed','expires_in':3600})) as refresh:
            self.assertEqual(google.token(),'renewed')
            self.assertEqual(refresh.call_args.kwargs['data']['grant_type'],'refresh_token')
        with patch('server.google.httpx.request',return_value=httpx.Response(200,json={'updatedRows':1})) as request:
            google.write('test-spreadsheet',"'Hoja'!A2",[['=SUM(1,2)','+literal','@literal',0]])
            self.assertEqual(request.call_args.kwargs['params']['valueInputOption'],'RAW')
            self.assertEqual(request.call_args.kwargs['json']['values'][0][0],'=SUM(1,2)')

    def test_retry_limit_and_pending_destination_survive_new_worker(self):
        from unittest.mock import patch
        destination=str(uuid4())
        other=str(uuid4())
        with connect() as db:
            db.execute('INSERT INTO destinations(id,spreadsheet,sheet_id,title) VALUES(?,?,0,?)',(destination,'first-'+destination,'Respuestas'))
            db.execute('INSERT INTO destinations(id,spreadsheet,sheet_id,title) VALUES(?,?,0,?)',(other,'second-'+other,'Otra'))
            db.execute('INSERT INTO integrations VALUES(?,?,1)',(self.fid,destination))
        payload=self.submission()
        self.client.post(f'/api/public/{self.fid}/submit',json=payload)
        with connect() as db:
            db.execute('UPDATE integrations SET destination_id=? WHERE form_id=?',(other,self.fid))
            db.execute('UPDATE jobs SET attempts=7 WHERE submission_id=?',(payload['id'],))
        fake=FakeSheets()
        with patch.object(fake,'metadata',side_effect=GoogleError('Sin permiso')):
            process_jobs(fake)
        with connect() as db:
            job=db.execute('SELECT * FROM jobs WHERE submission_id=?',(payload['id'],)).fetchone()
            self.assertEqual(job['attempts'],8)
            self.assertEqual(job['destination_id'],destination)
            db.execute('UPDATE jobs SET next_attempt=0 WHERE submission_id=?',(payload['id'],))
        process_jobs(FakeSheets())
        with connect() as db:
            self.assertEqual(db.execute('SELECT status FROM jobs WHERE submission_id=?',(payload['id'],)).fetchone()['status'],'error')
        self.client.post('/api/submissions/'+payload['id']+'/retry')
        fresh=FakeSheets()
        process_jobs(fresh)
        self.assertEqual(fresh.rows[2][0],payload['id'])


if __name__ == '__main__':
    unittest.main()
