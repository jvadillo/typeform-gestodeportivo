import io
import json
import re
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .engine import is_statement
from .variables import stored_variables


def make_workbook(submissions, documents):
    columns, variables = {}, {}
    for doc in documents:
        for q in doc['questions']:
            if not is_statement(q):
                columns.setdefault(q['id'], q['title'])
        for variable in doc.get('variables', []):
            variables[variable['id']] = variable['name']
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Respuestas'
    headers = ['Identificador del envío', 'Versión', 'Fecha y hora (UTC)', 'Puntuación', 'Resultado'] + [f'{title} [{qid}]' for qid, title in columns.items()]
    headers += [f'{name} [variable:{vid}]' for vid, name in variables.items()]
    if variables:
        headers[3] = 'Valor del resultado'
    # XML cannot represent these controls even when Excel cells are typed as strings.
    headers = [re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text) for text in headers]
    sheet.append(headers)
    for sub in submissions:
        answers = {a['question_id']: a['label'] for a in json.loads(sub['answers'])}
        date = datetime.fromisoformat(sub['completed']).astimezone(timezone.utc).replace(tzinfo=None)
        values = [sub['id'], sub['version'], date, sub['total'], sub['result_title']] + [answers.get(qid) for qid in columns]
        variable_values = {v['id']: v['value'] for v in stored_variables(sub)}
        values += [variable_values.get(vid) for vid in variables]
        values = [re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value) if isinstance(value, str) else value for value in values]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Explicit string, even for = + - @ participant input.
        sheet.cell(sheet.max_row, 3).number_format = 'yyyy-mm-dd hh:mm:ss'
    for cell in sheet[1]:
        cell.data_type = 's'
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='193D31')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 40
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 28
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
